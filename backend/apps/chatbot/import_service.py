"""
apps/chatbot/import_service.py
------------------------------
Service d'import multi-source vers la base de connaissances Sesily AI.

Sources supportées :
  pdf    — Documents PDF          (pdfplumber)
  excel  — Fichiers Excel / CSV   (pandas + openpyxl)
  docx   — Documents Word         (python-docx)
  url    — Pages web              (requests + beautifulsoup4)
  image  — Images avec texte      (Gemini Vision multimodal)
  text   — Texte brut             (direct)

Pipeline commun :
  extract_text() → structure_with_ai() → liste de dicts KnowledgeEntry
"""

import json
import logging
import re

import requests as http_requests
from django.conf import settings

logger = logging.getLogger(__name__)

# ── Dépendances optionnelles (dégradation gracieuse) ─────────────────────────
try:
    import pdfplumber          # pip install pdfplumber
    HAS_PDFPLUMBER = True
except ImportError:
    HAS_PDFPLUMBER = False

try:
    import docx as _docx      # pip install python-docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import pandas as pd        # déjà dans requirements.txt
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False

try:
    from bs4 import BeautifulSoup   # pip install beautifulsoup4
    HAS_BS4 = True
except ImportError:
    HAS_BS4 = False


# ── Prompt de structuration ───────────────────────────────────────────────────
_STRUCTURE_PROMPT = """\
Tu es un expert en agriculture à Madagascar travaillant pour SmartSaha.
Analyse le document suivant et extrait toutes les informations agronomiques exploitables.
Décompose-les en entrées distinctes de base de connaissances pour les agriculteurs malgaches.

Chaque entrée doit respecter ce format JSON exact :
{{
  "title":    "titre court et descriptif (50-100 caractères max)",
  "category": "exactement l'une de : maladie, calendrier, pratique, variete, marche, sol, meteo, stockage, elevage, general",
  "content":  "contenu complet en Markdown — utilise ## titres, **gras**, listes à puces, tableaux si pertinent",
  "crops":    ["liste des cultures : riz, maïs, manioc, haricot, patate douce, vanille, litchi, café, cacao, coton, arachide"],
  "region":   "région de Madagascar (Hautes Terres, Côte Est, Côte Ouest, Sud aride, Menabe, Itasy…) ou vide si tout Madagascar",
  "language": "fr, en, ou mg selon la langue du document"
}}

Règles :
- Une entrée = un sujet distinct (une maladie, une technique, un calendrier, une variété…)
- Le contenu doit être directement actionnable pour un agriculteur ou technicien
- Inclure les données chiffrées (prix, rendements, doses) si présentes dans le document
- Si aucune information agricole exploitable → retourne []
- Retourne UNIQUEMENT un tableau JSON valide, sans markdown, sans texte avant ou après

Document à analyser (source : {source_type}) :
---
{content}
---"""

VALID_CATEGORIES = frozenset({
    'maladie', 'calendrier', 'pratique', 'variete',
    'marche', 'sol', 'meteo', 'stockage', 'elevage', 'general'
})
VALID_LANGUAGES = frozenset({'fr', 'en', 'mg'})
MAX_CONTENT_CHARS = 60_000


class KBImportError(Exception):
    """Erreur levée lors d'un import de source."""


class ImportService:
    """Extrait et structure n'importe quelle source en liste de dicts KnowledgeEntry."""

    # ── Dispatch principal ────────────────────────────────────────────────────
    def extract_text(
        self,
        source_type: str,
        *,
        file=None,
        url: str | None = None,
        text: str | None = None,
    ) -> str:
        dispatch = {
            'pdf':   lambda: self._from_pdf(file),
            'excel': lambda: self._from_excel(file),
            'docx':  lambda: self._from_docx(file),
            'url':   lambda: self._from_url(url or ''),
            'image': lambda: self._from_image_gemini(file),
            'text':  lambda: (text or '').strip(),
        }
        fn = dispatch.get(source_type)
        if fn is None:
            raise KBImportError(f"Type de source inconnu : '{source_type}'")
        return fn()

    # ── Extracteurs ───────────────────────────────────────────────────────────
    def _from_pdf(self, file) -> str:
        if not HAS_PDFPLUMBER:
            raise KBImportError("pdfplumber non installé — exécutez : pip install pdfplumber")
        import pdfplumber as _pdf
        pages = []
        with _pdf.open(file) as pdf:
            for i, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text and page_text.strip():
                    pages.append(f"[Page {i + 1}]\n{page_text.strip()}")
        if not pages:
            raise KBImportError("Aucun texte extractible dans ce PDF (PDF scanné sans OCR ?)")
        return '\n\n'.join(pages)

    def _from_excel(self, file) -> str:
        if not HAS_PANDAS:
            raise KBImportError("pandas non installé — exécutez : pip install pandas openpyxl")
        parts = []
        try:
            sheets = pd.read_excel(file, sheet_name=None, engine='openpyxl')
            for sheet_name, df in sheets.items():
                df = df.dropna(how='all').fillna('')
                if df.empty:
                    continue
                parts.append(f"[Feuille : {sheet_name}]\n{df.to_string(index=False, max_rows=500)}")
        except Exception:
            # Fallback CSV
            try:
                file.seek(0)
                df = pd.read_csv(file, encoding='utf-8', errors='replace', nrows=1000)
                parts.append(df.to_string(index=False))
            except Exception as e:
                raise KBImportError(f"Impossible de lire ce fichier : {e}")
        if not parts:
            raise KBImportError("Fichier vide ou illisible.")
        return '\n\n'.join(parts)

    def _from_docx(self, file) -> str:
        if not HAS_DOCX:
            raise KBImportError("python-docx non installé — exécutez : pip install python-docx")
        doc = _docx.Document(file)
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = ' | '.join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    lines.append(row_text)
        if not lines:
            raise KBImportError("Aucun texte trouvé dans ce document Word.")
        return '\n'.join(lines)

    def _from_url(self, url: str) -> str:
        if not url.startswith(('http://', 'https://')):
            raise KBImportError("URL invalide (doit commencer par http:// ou https://)")
        try:
            resp = http_requests.get(
                url, timeout=20,
                headers={'User-Agent': 'SmartSaha-KB-Bot/1.0'},
                allow_redirects=True,
            )
            resp.raise_for_status()
        except http_requests.RequestException as e:
            raise KBImportError(f"Impossible d'accéder à l'URL : {e}")

        if HAS_BS4:
            soup = BeautifulSoup(resp.text, 'html.parser')
            for tag in soup(['script', 'style', 'nav', 'footer', 'header', 'aside', 'form']):
                tag.decompose()
            text = soup.get_text(separator='\n', strip=True)
        else:
            text = re.sub(r'<[^>]+>', ' ', resp.text)
            text = re.sub(r'\s+', ' ', text).strip()

        if len(text) < 80:
            raise KBImportError("Page trop courte ou contenu non extractible.")
        return text

    def _from_image_gemini(self, file) -> str:
        """Gemini Vision lit l'image directement (multimodal)."""
        api_key = getattr(settings, 'GEMINI_API_KEY', None)
        if not api_key:
            raise KBImportError("GEMINI_API_KEY non configuré dans les settings.")
        try:
            from google import genai
            from google.genai import types as genai_types
        except ImportError:
            raise KBImportError("google-genai non installé — exécutez : pip install google-genai")

        image_bytes = file.read()
        mime = getattr(file, 'content_type', None) or 'image/jpeg'
        client = genai.Client(api_key=api_key)
        response = client.models.generate_content(
            model='gemini-2.0-flash',
            contents=[
                genai_types.Part.from_bytes(data=image_bytes, mime_type=mime),
                "Extrais tout le texte agricole de cette image. Retourne uniquement le texte brut.",
            ],
        )
        extracted = response.text.strip()
        if not extracted:
            raise KBImportError("Aucun texte extrait de l'image.")
        return extracted

    # ── Structuration IA ──────────────────────────────────────────────────────
    def structure_with_ai(self, raw_content: str, source_type: str = 'document') -> list[dict]:
        """Envoie le texte brut à Gemini → liste de dicts prêts pour KnowledgeEntry."""
        api_key = getattr(settings, 'GEMINI_API_KEY', None)
        if not api_key:
            raise KBImportError("GEMINI_API_KEY non configuré.")
        try:
            from google import genai
        except ImportError:
            raise KBImportError("google-genai non installé.")

        content = raw_content[:MAX_CONTENT_CHARS]
        prompt = _STRUCTURE_PROMPT.format(source_type=source_type, content=content)

        client = genai.Client(api_key=api_key)
        response = client.models.generate_content(model='gemini-2.0-flash', contents=prompt)

        raw = response.text.strip()
        # Supprime les blocs markdown éventuels
        raw = re.sub(r'^```(?:json)?\s*\n?', '', raw, flags=re.MULTILINE)
        raw = re.sub(r'\n?```\s*$', '', raw, flags=re.MULTILINE).strip()

        try:
            data = json.loads(raw)
        except json.JSONDecodeError:
            # Tentative de récupération — cherche le premier tableau JSON
            match = re.search(r'\[.*\]', raw, re.DOTALL)
            if match:
                try:
                    data = json.loads(match.group())
                except json.JSONDecodeError:
                    raise KBImportError("L'IA n'a pas retourné un JSON valide. Réessayez.")
            else:
                raise KBImportError("L'IA n'a pas retourné un JSON valide. Réessayez.")

        if not isinstance(data, list):
            raise KBImportError("La réponse de l'IA n'est pas un tableau JSON.")

        return [self._clean_entry(item) for item in data if self._is_valid_entry(item)]

    # ── Helpers ───────────────────────────────────────────────────────────────
    @staticmethod
    def _is_valid_entry(item: dict) -> bool:
        return (
            isinstance(item, dict)
            and bool(str(item.get('title', '')).strip())
            and bool(str(item.get('content', '')).strip())
        )

    @staticmethod
    def _clean_entry(item: dict) -> dict:
        title    = str(item.get('title', '')).strip()[:150]
        content  = str(item.get('content', '')).strip()
        category = str(item.get('category', 'general')).lower().strip()
        language = str(item.get('language', 'fr')).lower().strip()
        crops_raw = item.get('crops', [])
        crops = [str(c).strip().lower() for c in (crops_raw if isinstance(crops_raw, list) else []) if str(c).strip()]
        return {
            'title':    title,
            'category': category if category in VALID_CATEGORIES else 'general',
            'content':  content,
            'crops':    crops,
            'region':   str(item.get('region', '')).strip(),
            'language': language if language in VALID_LANGUAGES else 'fr',
        }
